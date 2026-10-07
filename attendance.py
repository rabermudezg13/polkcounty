"""Workbook parsing and attendance rules. VBA macros are never executed."""
import csv
import hashlib
import io
import re
from collections import Counter
from datetime import date, datetime

import openpyxl
from openpyxl.utils.datetime import from_excel

STATUSES = ('No Show', 'Late Cancellation')
SOURCE_SHEETS = ('Occurrence Log', 'Assisted', 'Unassisted', 'No Show Import')


def clean(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_status(value):
    text = re.sub(r'[\s_-]+', ' ', clean(value).casefold())
    return {'no show': 'No Show', 'noshow': 'No Show', 'ncns': 'No Show',
            'no call no show': 'No Show', 'late cancellation': 'Late Cancellation',
            'unassisted': 'Late Cancellation'}.get(text)


def is_polk(value):
    text = clean(value).casefold()
    return text.split('|')[-1].strip() in ('polk', 'polk county', 'polk county schools')


def person_id(ats_id, email, name):
    if clean(ats_id):
        return 'ats:' + clean(ats_id)
    # The workbook has missing ATS IDs; names are the consistent fallback
    # shared by assisted and unassisted reports (which have no email).
    return 'name:' + ' '.join(clean(name).casefold().split())


def occurrence_id(row):
    identity = '\x1f'.join([row['person_id'], row['date'], row['status']])
    return hashlib.sha256(identity.encode()).hexdigest()


def parse_date(value, epoch):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            result = from_excel(value, epoch)
            return result.date().isoformat()
        except (ValueError, OverflowError, AttributeError):
            return None
    text = clean(value)
    for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%m/%d/%Y', '%m/%d/%y',
                '%m/%d/%Y %H:%M', '%m/%d/%Y %I:%M %p'):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def sheet_names(data):
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [name for name in SOURCE_SHEETS if name in book.sheetnames]
    finally:
        book.close()


def parse_workbook(data, sheet='Occurrence Log'):
    """Return Polk records plus an explicit account of excluded/invalid rows."""
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    counts = Counter()
    errors, records = [], {}
    try:
        if sheet not in SOURCE_SHEETS or sheet not in book.sheetnames:
            raise ValueError('Select a supported import sheet.')
        rows = list(book[sheet].values)
        header_index = None
        for i, values in enumerate(rows[:15]):
            labels = {clean(v) for v in values}
            if ('Employee Name' in labels and 'Date of Absence' in labels) or \
               ('Talent Name' in labels and 'Date of NCNS' in labels) or \
               ('Cancelling Sub Name' in labels and 'Absence Date' in labels):
                header_index = i
                break
        if header_index is None:
            raise ValueError('Expected attendance headers were not found.')
        headers = [clean(v) for v in rows[header_index]]
        for number, values in enumerate(rows[header_index + 1:], header_index + 2):
            if not any(clean(v) for v in values):
                continue
            row = dict(zip(headers, values))
            if clean(values[0]).startswith('←') or row.get('District') == 'District':
                counts['Template/header rows skipped'] += 1
                continue
            district = row.get('District', row.get('Org Name', ''))
            if not is_polk(district):
                counts['Other districts excluded'] += 1
                continue
            name = clean(row.get('Employee Name', row.get('Talent Name', row.get('Cancelling Sub Name'))))
            raw_type = row.get('Type') if sheet == 'Occurrence Log' else \
                ('Unassisted' if sheet == 'Unassisted' else row.get('Subcategory'))
            if sheet == 'No Show Import' and not clean(raw_type):
                raw_type = 'No Show'
            status = normalize_status(raw_type)
            day = parse_date(row.get('Date of Absence', row.get('Date of NCNS', row.get('Absence Date'))), book.epoch)
            if not name or not day or not status:
                errors.append({'Sheet': sheet, 'Row': number,
                    'Issue': 'Missing name, invalid date or unsupported incident type', 'Type': clean(raw_type)})
                continue
            ats = clean(row.get('ATS ID', row.get('KSN ID')))
            email = clean(row.get('Email'))
            record = {'person_id': person_id(ats, email, name), 'ats_id': ats, 'name': name,
                'email': email, 'date': day, 'status': status, 'district': 'Polk',
                'school': clean(row.get('School', row.get('School Information', row.get('Location Name')))),
                'case_number': clean(row.get('Case Number', row.get('Confirmation #'))),
                'source_type': clean(raw_type), 'source_sheet': sheet,
                'notes': clean(row.get('Notes'))}
            key = occurrence_id(record)
            record['id'] = key
            if key in records:
                counts['Duplicate incidents excluded'] += 1
                continue
            records[key] = record
        counts['Valid Polk incidents'] = len(records)
        counts['Rows requiring correction'] = len(errors)
        counts['Name-based identities'] = sum(not r['ats_id'] for r in records.values())
        return list(records.values()), dict(counts), errors
    finally:
        book.close()


def recurring_people(records, threshold=2):
    people = {}
    # Stable document IDs make repeats safe; handle duplicated input defensively.
    for row in {occurrence_id(r): r for r in records}.values():
        p = people.setdefault(row['person_id'], {'Person ID': row['person_id'],
            'Name': row['name'], 'ATS ID': row.get('ats_id', ''),
            'Email': row.get('email', ''), 'No Show': 0, 'Late Cancellation': 0,
            'Last incident': row['date']})
        if row['status'] in STATUSES:
            p[row['status']] += 1
            p['Last incident'] = max(p['Last incident'], row['date'])
    return sorted([dict(p, Incidents=p['No Show'] + p['Late Cancellation']) for p in people.values()
                   if p['No Show'] + p['Late Cancellation'] >= threshold],
                  key=lambda p: (-p['Incidents'], p['Name'].casefold()))


def export_csv(records):
    """Make spreadsheet downloads safe from formula injection."""
    if not records:
        return b''
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=list(records[0]))
    writer.writeheader()
    for row in records:
        writer.writerow({k: "'" + v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@'))
                         else v for k, v in row.items()})
    return output.getvalue().encode('utf-8-sig')
